# -*- coding: utf-8 -*-
import logging
import re
from datetime import datetime, time

from markupsafe import Markup, escape

from odoo import Command, _, api, fields, models
from odoo.exceptions import UserError

from odoo.addons.tectora_products.models import catalog_rules

from ..models import quote_rules

_logger = logging.getLogger(__name__)

GENERIC_PRODUCT_CODE = "TECTORA-IMPORT"
# Units of the old system that catalog_rules.normalise_uom does not know.
EXTRA_UOM_KEYS = {"uren": "hour", "manuur": "hour", "dagen": "day"}
# Context of every write of the import: no mails, no chatter tracking.
QUIET = {"tracking_disable": True, "mail_create_nolog": True,
         "mail_notrack": True, "mail_auto_subscribe_no_notify": True}


def _norm(text):
    return re.sub(r"\s+", " ", (text or "").strip()).lower()


class _DryRun(Exception):
    """Raised to roll a test run back."""


class QuotationImport(models.TransientModel):
    _name = "tectora.quotation.import"
    _description = "Offertes importeren uit het oude systeem"

    file = fields.Binary(string="Bestand (.xlsx)", required=True, attachment=False)
    filename = fields.Char(string="Bestandsnaam")
    sheet_name = fields.Char(
        string="Tabblad",
        help="Laat leeg om het eerste tabblad (of 'Export Simpla') te nemen.",
    )
    dry_run = fields.Boolean(
        string="Testrun (niets opslaan)",
        help="Leest en verwerkt het hele bestand en toont het resultaat, maar "
        "draait alles terug: handig om een nieuwe lijst eerst na te kijken.",
    )
    existing_mode = fields.Selection(
        [
            ("update", "Status bijwerken"),
            ("skip", "Overslaan"),
        ],
        string="Gekende offertes",
        default="update",
        required=True,
        help="Een offerte waarvan het nummer al in Odoo staat, wordt nooit een "
        "tweede keer aangemaakt. 'Status bijwerken' brengt een offerte die "
        "intussen goedgekeurd werd naar een bevestigde order (en de "
        "verkoopkans naar gewonnen); orderlijnen worden niet aangeraakt en een "
        "bevestigde order wordt nooit teruggezet.",
    )
    use_legacy_name = fields.Boolean(
        string="Offertenummer als ordernummer",
        default=True,
        help="De order krijgt het nummer uit het oude systeem (bv. 2026/0004) "
        "in plaats van een nieuw Odoo-nummer. Het oude nummer staat hoe dan "
        "ook op de order (Offertenummer oud systeem).",
    )
    user_id = fields.Many2one(
        "res.users", string="Verkoper", default=lambda self: self.env.user,
        domain="[('share', '=', False)]",
    )
    team_id = fields.Many2one(
        "crm.team", string="Verkoopteam",
        default=lambda self: self.env["crm.team"]._get_default_team_id(),
    )
    tax_0_id = fields.Many2one(
        "account.tax", string="Btw 0%",
        default=lambda self: self._find_tax(0.0),
        domain="[('type_tax_use', '=', 'sale'), ('amount', '=', 0)]",
        help="Btw voor de lijnen aan 0% (standaard medecontractant).",
    )
    tax_6_id = fields.Many2one(
        "account.tax", string="Btw 6%",
        default=lambda self: self._find_tax(6.0),
        domain="[('type_tax_use', '=', 'sale'), ('amount', '=', 6)]",
    )
    tax_21_id = fields.Many2one(
        "account.tax", string="Btw 21%",
        default=lambda self: self._find_tax(21.0),
        domain="[('type_tax_use', '=', 'sale'), ('amount', '=', 21)]",
    )
    text_lines = fields.Selection(
        [
            ("skip", "Overslaan"),
            ("note", "Als notitie onderaan de offerte"),
        ],
        string="Titel- en tekstregels",
        default="skip",
        required=True,
        help="Regels zonder aantal en prijs (hoofdstuktitels zoals ALGEMENE "
        "WERKEN, nota's). De export levert de lijnen alfabetisch, dus een "
        "titel staat niet meer boven zijn lijnen.",
    )
    limit = fields.Integer(
        string="Maximum aantal nieuwe offertes",
        help="0 = alles. Bij een grote lijst kan je in delen importeren: "
        "start de import opnieuw met hetzelfde bestand en de reeds "
        "geïmporteerde offertes worden overgeslagen.",
    )
    state = fields.Selection(
        [("draft", "Concept"), ("done", "Uitgevoerd")], default="draft"
    )
    summary = fields.Text(string="Resultaat", readonly=True)

    # ------------------------------------------------------------------ setup
    @api.model
    def _find_tax(self, rate):
        """The sale tax for a rate: for works the service variant ('6% S'),
        at 0% the co-contractor tax (B2B works in Belgium)."""
        Tax = self.env["account.tax"]
        company = self.env.company
        taxes = Tax.search(
            [
                ("type_tax_use", "=", "sale"),
                ("amount_type", "=", "percent"),
                ("amount", "=", rate),
                ("company_id", "in", company.parent_ids.ids),
            ],
            order="sequence, id",
        ).filtered(lambda tax: not tax.price_include)
        preferred = ["cocont", "medecontract"] if not rate else ["%g%% s" % rate]
        for word in preferred:
            hit = taxes.filtered(lambda tax: word in (tax.name or "").lower())
            if hit:
                return hit[:1]
        return taxes[:1]

    def _read_rows(self):
        """Return (header, rows) of the uploaded workbook."""
        self.ensure_one()
        try:
            import openpyxl
        except ImportError:
            raise UserError(
                _("De Python-bibliotheek openpyxl is niet beschikbaar op de server.")
            )
        try:
            workbook = openpyxl.load_workbook(
                self.file.open(), data_only=True, read_only=True
            )
        except Exception as error:
            raise UserError(
                _("Het bestand kon niet gelezen worden als .xlsx-bestand (%s).", error)
            )
        name = (self.sheet_name or "").strip()
        if name and name in workbook.sheetnames:
            sheet = workbook[name]
        elif name:
            raise UserError(
                _("Tabblad '%(sheet)s' niet gevonden. Beschikbaar: %(all)s",
                  sheet=name, all=", ".join(workbook.sheetnames))
            )
        elif "Export Simpla" in workbook.sheetnames:
            sheet = workbook["Export Simpla"]
        else:
            sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            raise UserError(_("Het tabblad bevat geen gegevens."))
        return rows[0], rows[1:]

    def _statuses(self, quotes):
        """{status key: tectora.quotation.import.status}; every status of the
        file must be configured before anything is written."""
        Status = self.env["tectora.quotation.import.status"]
        statuses = {Status._key(status.name): status for status in Status.search([])}
        unknown = sorted({
            quote["status"] or _("(leeg)") for quote in quotes
            if Status._key(quote["status"]) not in statuses
        })
        if unknown:
            raise UserError(_(
                "Onbekende status in het bestand: %(statuses)s.\n\nVoeg ze toe "
                "in Verkoop → Configuratie → Statussen oud systeem (welke "
                "orderstatus en fase van de verkoopkans ze krijgen) en start "
                "de import opnieuw.",
                statuses=", ".join(unknown),
            ))
        return statuses

    # --------------------------------------------------------------- lookups
    def _prepare_lookups(self):
        """Works items by name, partners by VAT number and by name + postcode."""
        company = self.env.company
        # Only works items (services): a material on an imported order would
        # make a delivery when the order is confirmed.
        products = {}
        for product in self.env["product.product"].search(
            [
                ("sale_ok", "=", True),
                ("type", "=", "service"),
                ("company_id", "in", [company.id, False]),
            ],
            order="id",
        ):
            products.setdefault(_norm(product.name), product)
        partners_by_vat = {}
        partners_by_name = {}
        for partner in self.env["res.partner"].search(
            [("company_id", "in", [company.id, False]), ("parent_id", "=", False)],
            order="id",
        ):
            vat = quote_rules.normalise_vat(partner.vat)
            if vat:
                partners_by_vat.setdefault(vat, partner)
            partners_by_name.setdefault((_norm(partner.name), (partner.zip or "").strip()), partner)
        return {
            "products": products,
            "partners_by_vat": partners_by_vat,
            "partners_by_name": partners_by_name,
            "uoms": {},
            "taxes": {0.0: self.tax_0_id, 6.0: self.tax_6_id, 21.0: self.tax_21_id},
            "generic": None,
        }

    def _generic_product(self, lookups):
        """The works item for lines of the old system without a product of
        their own in the catalogue: the line keeps its description, unit and
        price."""
        if lookups["generic"]:
            return lookups["generic"]
        Product = self.env["product.product"].with_context(active_test=False)
        product = Product.search([("default_code", "=", GENERIC_PRODUCT_CODE)], limit=1)
        if not product:
            product = Product.create({
                "name": _("Werkpost (oud systeem)"),
                "default_code": GENERIC_PRODUCT_CODE,
                "type": "service",
                "sale_ok": True,
                "purchase_ok": False,
                "invoice_policy": "order",
                "list_price": 0.0,
                "taxes_id": [Command.clear()],
                "description_sale": _(
                    "Lijn uit een offerte van het oude systeem zonder "
                    "overeenkomstig product in de catalogus."
                ),
            })
        elif not product.active:
            product.active = True
        lookups["generic"] = product
        return product

    def _uom(self, unit, lookups):
        """The unit of measure for a unit of the file (None: the product's)."""
        key = _norm(unit)
        if not key:
            return None
        if key in lookups["uoms"]:
            return lookups["uoms"][key]
        uom = None
        uom_key = EXTRA_UOM_KEYS.get(key) or catalog_rules.normalise_uom(key)
        if uom_key != "unit" or key in catalog_rules.PLAIN_UNIT_WORDS:
            uom = self.env["product.template"]._tectora_get_uom(
                uom_key, catalog_rules.UOM_NAMES.get(uom_key, "Stuk")
            )
        else:
            # km, m³, week, ...: an existing unit of that name, else the product's.
            uom = self.env["uom.uom"].search([("name", "=ilike", unit.strip())], limit=1) or None
        lookups["uoms"][key] = uom
        return uom

    def _tax(self, rate, lookups):
        rate = round(rate, 2)
        if rate not in lookups["taxes"] or not lookups["taxes"][rate]:
            lookups["taxes"][rate] = self._find_tax(rate)
        return lookups["taxes"][rate]

    def _partner(self, customer, lookups, counters):
        """The customer: matched on VAT number, then on name and postcode;
        created when unknown."""
        vat = customer["vat"]
        if vat and vat in lookups["partners_by_vat"]:
            return lookups["partners_by_vat"][vat]
        name_key = (_norm(customer["name"]), customer["zip"])
        if name_key in lookups["partners_by_name"]:
            partner = lookups["partners_by_name"][name_key]
            if vat:
                lookups["partners_by_vat"].setdefault(vat, partner)
            return partner
        Partner = self.env["res.partner"]
        belgium = self.env.ref("base.be", raise_if_not_found=False)
        values = {
            "name": customer["name"] or _("Klant zonder naam"),
            "street": customer["street"] or False,
            "zip": customer["zip"] or False,
            "city": customer["city"] or False,
            "country_id": belgium.id if belgium else False,
            "customer_rank": 1,
        }
        if vat:
            checked, _country = Partner._run_vat_checks(
                belgium, vat, partner_name=values["name"], validation="setnull"
            )
            if checked:
                values["vat"] = checked
            else:
                values["comment"] = _("Btw-nummer uit het oude systeem (ongeldig): %s",
                                      customer["vat_raw"])
                counters["invalid_vat"].append(customer["vat_raw"])
        partner = Partner.with_context(**QUIET).create(values)
        counters["partners"] += 1
        if vat:
            lookups["partners_by_vat"][vat] = partner
        lookups["partners_by_name"][name_key] = partner
        return partner

    # ------------------------------------------------------------ the orders
    def _line_values(self, line, sequence, lookups, counters):
        if line["kind"] == "text":
            return {
                "display_type": "line_note",
                "name": line["description"],
                "sequence": sequence,
            }
        product = lookups["products"].get(_norm(line["description"]))
        if product:
            counters["matched_lines"] += 1
        else:
            product = self._generic_product(lookups)
            counters["generic_lines"] += 1
        uom = self._uom(line["unit"], lookups)
        tax = self._tax(line["vat_rate"], lookups)
        if not tax:
            counters["missing_taxes"].add(line["vat_rate"])
        price = line["price"]
        return {
            "product_id": product.id,
            "name": line["description"] or product.display_name,
            "product_uom_qty": line["quantity"],
            "product_uom_id": (uom or product.uom_id).id,
            "price_unit": price,
            # The price is the one of the old quotation. Odoo treats a price
            # that differs from technical_price_unit as set by hand and then
            # never re-prices the line from the pricelist (on a quantity
            # change, ...); the generic works item would otherwise drop to 0.
            "technical_price_unit": price + 1.0,
            "discount": line["discount"],
            "tax_ids": [Command.set(tax.ids)],
            "sequence": sequence,
        }

    def _order_values(self, quote, partner, lead, status, lookups, counters):
        lines = []
        sequence = 10
        for line in quote["lines"]:
            if line["kind"] == "text":
                continue
            lines.append(Command.create(self._line_values(line, sequence, lookups, counters)))
            sequence += 10
        texts = [line for line in quote["lines"] if line["kind"] == "text"]
        if texts and self.text_lines == "note":
            lines.append(Command.create({
                "display_type": "line_section",
                "name": _("Titels en nota's uit het oude systeem"),
                "sequence": sequence,
            }))
            for line in texts:
                sequence += 10
                lines.append(Command.create(self._line_values(line, sequence, lookups, counters)))
        counters["skipped_text"] += 0 if self.text_lines == "note" else len(texts)
        values = {
            "partner_id": partner.id,
            "company_id": self.env.company.id,
            "user_id": self.user_id.id or False,
            "team_id": self.team_id.id or False,
            "opportunity_id": lead.id,
            "client_order_ref": quote["customer_ref"] or False,
            "origin": _("Oud systeem %s", quote["number"]),
            "tectora_legacy_number": quote["number"],
            "tectora_legacy_status": quote["status"],
            "tectora_legacy_invoiced": status.invoiced_legacy,
            "order_line": lines,
        }
        if quote["date"]:
            values["date_order"] = self._datetime(quote["date"])
        if quote["validity_date"]:
            values["validity_date"] = quote["validity_date"]
        if self.use_legacy_name:
            values["name"] = quote["number"]
        return values

    @api.model
    def _datetime(self, day):
        # Noon, so the date stays the same in every timezone.
        return datetime.combine(day, time(12, 0))

    def _lead_values(self, quote, partner, status):
        customer = quote["customer"]
        title = quote["customer_ref"] or customer["name"]
        values = {
            "name": "%s - %s" % (quote["number"], title) if title else quote["number"],
            "type": "opportunity",
            "partner_id": partner.id,
            "user_id": self.user_id.id or False,
            "team_id": self.team_id.id or False,
            "company_id": self.env.company.id,
            "expected_revenue": quote["total"],
            "street": customer["street"] or False,
            "zip": customer["zip"] or False,
            "city": customer["city"] or False,
            "stage_id": status._resolve_stage(self.team_id).id,
        }
        if status.lead_outcome == "won":
            values["probability"] = 100.0
        return values

    def _apply_status(self, order, status, quote):
        """Bring the order and its opportunity in the state of the status."""
        lead = order.opportunity_id
        quiet = order.with_context(**QUIET)
        if not status.create_roof_project:
            quiet = quiet.with_context(tectora_import_no_project=True)
        target = status.order_state
        if target == "sent" and order.state == "draft":
            quiet.write({"state": "sent"})
        elif target == "sale" and order.state in ("draft", "sent"):
            quiet.action_confirm()
            if quote["date"]:
                # Confirming dates the order today; it was signed back then.
                quiet.write({"date_order": self._datetime(quote["date"])})
        elif target == "cancel" and order.state in ("draft", "sent"):
            quiet._action_cancel()
        if lead:
            lead = lead.with_context(**QUIET)
            stage = status._resolve_stage(self.team_id)
            if status.lead_outcome == "lost":
                if lead.active:
                    lead.action_set_lost()
            elif stage and lead.stage_id != stage:
                values = {"stage_id": stage.id}
                if status.lead_outcome == "won":
                    values["probability"] = 100.0
                lead.write(values)
            if status.lead_outcome == "won" and quote["date"]:
                lead.write({"date_closed": self._datetime(quote["date"])})

    def _post_notes(self, order, quote):
        if not quote["notes"]:
            return
        body = Markup("<p><b>%s</b></p><p>%s</p>") % (
            _("Opmerkingen uit het oude systeem"),
            Markup("<br/>").join(escape(line) for line in quote["notes"].splitlines()),
        )
        order.with_context(**QUIET).message_post(body=body, message_type="comment",
                                                 subtype_xmlid="mail.mt_note")

    def _import_quote(self, quote, status, lookups, counters):
        partner = self._partner(quote["customer"], lookups, counters)
        lead = self.env["crm.lead"].with_context(**QUIET).create(
            self._lead_values(quote, partner, status)
        )
        Order = self.env["sale.order"].with_context(**QUIET)
        if not status.create_roof_project:
            Order = Order.with_context(tectora_no_roof_project=True)
        order = Order.create(self._order_values(quote, partner, lead, status, lookups, counters))
        self._apply_status(order, status, quote)
        self._post_notes(order, quote)
        if order.currency_id.compare_amounts(order.amount_untaxed, quote["total"]):
            counters["total_differences"].append(
                "%s (%.2f / %.2f)" % (quote["number"], order.amount_untaxed, quote["total"])
            )
        return order

    def _update_quote(self, order, quote, status, counters):
        """A quotation already in Odoo: only move it forward."""
        if order.tectora_legacy_status == quote["status"]:
            counters["unchanged"] += 1
            return
        if order.state not in ("draft", "sent") and not (
            order.state == "sale" and status.order_state == "sale"
        ):
            counters["not_reverted"].append(
                "%s (%s → %s)" % (quote["number"], order.tectora_legacy_status, quote["status"])
            )
            return
        self._apply_status(order, status, quote)
        order.with_context(**QUIET).write({
            "tectora_legacy_status": quote["status"],
            "tectora_legacy_invoiced": status.invoiced_legacy,
        })
        counters["updated"] += 1

    # ------------------------------------------------------------- the run
    def action_import(self):
        self.ensure_one()
        header, rows = self._read_rows()
        try:
            quotes, stats = quote_rules.parse_rows(header, rows)
        except ValueError as error:
            raise UserError(str(error))
        if not quotes:
            raise UserError(_("Geen offertes gevonden in het bestand."))
        if stats["bad_dates"]:
            raise UserError(_("Offertes zonder leesbare offertedatum: %s",
                              ", ".join(stats["bad_dates"][:20])))
        statuses = self._statuses(quotes)

        counters = {
            "created": 0, "updated": 0, "unchanged": 0, "skipped": 0,
            "postponed": 0, "partners": 0, "matched_lines": 0,
            "generic_lines": 0, "skipped_text": 0, "invalid_vat": [],
            "missing_taxes": set(), "total_differences": [], "not_reverted": [],
            "errors": [], "per_status": {},
        }
        try:
            with self.env.cr.savepoint():
                self._run(quotes, statuses, counters)
                if self.dry_run:
                    raise _DryRun()  # rolls the savepoint back
        except _DryRun:
            pass
        _logger.info("tectora_quotation_import: %s (%s)", self.filename, {
            key: value for key, value in counters.items() if isinstance(value, int)
        })
        self.write({"state": "done", "summary": self._summary(quotes, stats, counters)})
        return {
            "type": "ir.actions.act_window",
            "res_model": self._name,
            "res_id": self.id,
            "view_mode": "form",
            "target": "new",
            "context": self.env.context,
        }

    def _run(self, quotes, statuses, counters):
        lookups = self._prepare_lookups()
        Order = self.env["sale.order"].with_context(active_test=False)
        existing = {
            order.tectora_legacy_number: order
            for order in Order.search([
                ("tectora_legacy_number", "in", [quote["number"] for quote in quotes]),
                ("company_id", "=", self.env.company.id),
            ])
        }
        for quote in quotes:
            status = statuses[self.env["tectora.quotation.import.status"]._key(quote["status"])]
            order = existing.get(quote["number"])
            try:
                with self.env.cr.savepoint():
                    if order:
                        if self.existing_mode == "update":
                            self._update_quote(order, quote, status, counters)
                        else:
                            counters["skipped"] += 1
                        continue
                    if self.limit and counters["created"] >= self.limit:
                        counters["postponed"] += 1
                        continue
                    self._import_quote(quote, status, lookups, counters)
                    counters["created"] += 1
                    counters["per_status"][quote["status"]] = (
                        counters["per_status"].get(quote["status"], 0) + 1
                    )
            except Exception as error:
                _logger.exception("tectora_quotation_import: offerte %s", quote["number"])
                counters["errors"].append("%s: %s" % (quote["number"], error))
                # Partners, uoms and the generic product made inside the
                # rolled back savepoint are gone: look them up again.
                lookups = self._prepare_lookups()

    def _summary(self, quotes, stats, counters):
        lines = []
        if self.dry_run:
            lines += [_("TESTRUN: er werd niets opgeslagen."), ""]
        lines.append(_("Bestand: %(file)s — %(quotes)s offertes, %(rows)s lijnen",
                       file=self.filename or "-", quotes=len(quotes), rows=stats["rows"]))
        lines.append(_("%s offertes aangemaakt (met verkoopkans)", counters["created"]))
        for status, count in sorted(counters["per_status"].items()):
            lines.append("    %s: %s" % (status, count))
        if counters["updated"] or self.existing_mode == "update":
            lines.append(_("%(upd)s gekende offertes bijgewerkt, %(same)s ongewijzigd",
                           upd=counters["updated"], same=counters["unchanged"]))
        if counters["skipped"]:
            lines.append(_("%s gekende offertes overgeslagen", counters["skipped"]))
        if counters["postponed"]:
            lines.append(_("%s offertes nog niet geïmporteerd (maximum bereikt): "
                           "start de import opnieuw met hetzelfde bestand",
                           counters["postponed"]))
        lines.append(_("%s nieuwe klanten aangemaakt", counters["partners"]))
        lines.append(_("%(matched)s lijnen gekoppeld aan een product uit de "
                       "catalogus, %(generic)s op de werkpost 'oud systeem' "
                       "(met hun eigen omschrijving, eenheid en prijs)",
                       matched=counters["matched_lines"], generic=counters["generic_lines"]))
        if stats["multiplied"]:
            lines.append(_("%s lijnen waarvan het bedrag een veelvoud was van aantal × "
                           "prijs: het aantal werd vermenigvuldigd zodat het bedrag "
                           "klopt", stats["multiplied"]))
        if stats["repriced"]:
            lines.append(_("%s lijnen waarvan de eenheidsprijs uit het bedrag "
                           "werd herrekend", stats["repriced"]))
        if counters["skipped_text"]:
            lines.append(_("%s titel- en tekstregels overgeslagen", counters["skipped_text"]))
        if counters["invalid_vat"]:
            lines.append(_("Ongeldige btw-nummers (in de notitie van de klant gezet): %s",
                           ", ".join(sorted(set(counters["invalid_vat"])))))
        if counters["missing_taxes"]:
            lines.append(_("Geen verkoopbtw gevonden voor: %s — die lijnen staan zonder btw",
                           ", ".join("%g%%" % rate for rate in sorted(counters["missing_taxes"]))))
        if counters["total_differences"]:
            lines.append(_("Totaal excl. btw wijkt af van het bestand (Odoo / bestand): %s",
                           ", ".join(counters["total_differences"])))
        if counters["not_reverted"]:
            lines.append(_("Niet teruggezet (een order gaat nooit achteruit): %s",
                           ", ".join(counters["not_reverted"])))
        if stats["conflicts"]:
            lines.append(_("Offertes met verschillende klant of status op hun "
                           "lijnen (eerste lijn genomen): %s", ", ".join(stats["conflicts"])))
        if stats["skipped_rows"]:
            lines.append(_("%s rijen zonder offertenummer overgeslagen", stats["skipped_rows"]))
        if counters["errors"]:
            lines += ["", _("%s offertes NIET geïmporteerd door een fout:",
                            len(counters["errors"]))]
            lines += ["    " + error for error in counters["errors"]]
        return "\n".join(lines)

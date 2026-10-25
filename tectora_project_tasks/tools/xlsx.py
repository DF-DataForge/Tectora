# -*- coding: utf-8 -*-
"""Reading and writing the standardised Excel lists of the project follow-up
(task templates, order list): one tab per list, a header row with fixed
column names, one record per row."""
import io
import re
import unicodedata

TRUE_WORDS = {"ja", "j", "x", "1", "yes", "y", "true", "waar", "optie", "v"}
FALSE_WORDS = {"", "nee", "n", "0", "no", "false", "onwaar"}


def normalize(text):
    """Header and keyword matching: lowercase, no accents, single spaces."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", text.lower())).strip()


def cell_text(value):
    """A cell as text: numbers without a needless ".0", None as ""."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_bool(value):
    """Yes/no from a cell; None when the value is not understood."""
    if isinstance(value, bool):
        return value
    word = normalize(cell_text(value))
    if word in TRUE_WORDS:
        return True
    if word in FALSE_WORDS:
        return False
    return None


def parse_int(value):
    """A whole number from a cell ("-7", 14.0); None when empty, ValueError
    when it is not a number."""
    text = cell_text(value).replace(",", ".")
    if not text:
        return None
    number = float(text)
    if not number.is_integer():
        raise ValueError(text)
    return int(number)


def read_sheet(data, sheet_names, columns, required_columns=(), required=True):
    """Rows of the first tab named in `sheet_names` (or the first tab when
    `required` and none is found) as dicts keyed by the keys of `columns`
    ({key: (header, aliases...)}), each with its Excel row number under
    "_row". Empty rows are skipped. Returns None when the tab is optional
    and absent. Raises ValueError for an unreadable file or when one of
    `required_columns` (keys) is missing.
    """
    import openpyxl  # shipped with Odoo (requirements.txt)

    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as error:
        raise ValueError(f"Het bestand kon niet gelezen worden als .xlsx-bestand ({error}).") from None
    by_name = {normalize(name): name for name in workbook.sheetnames}
    sheet = next((workbook[by_name[normalize(n)]] for n in sheet_names if normalize(n) in by_name), None)
    if sheet is None:
        if not required:
            return None
        sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None) or ()
    positions = {normalize(title): index for index, title in enumerate(header) if cell_text(title)}
    index_of = {}
    for key, names in columns.items():
        index = next((positions[normalize(n)] for n in names if normalize(n) in positions), None)
        if index is not None:
            index_of[key] = index
    missing = [columns[key][0] for key in required_columns if key not in index_of]
    if missing:
        raise ValueError(
            f"Tabblad '{sheet.title}': verplichte kolom(men) ontbreken: " + ", ".join(missing)
        )
    result = []
    for number, row in enumerate(rows, start=2):
        values = {key: (row[index] if index < len(row) else None) for key, index in index_of.items()}
        if not any(cell_text(value) for value in values.values()):
            continue
        values["_row"] = number
        result.append(values)
    return result


def build_workbook(sheets):
    """An .xlsx file (bytes) from [(title, header, rows, widths)]: bold
    header row, frozen, column widths in characters."""
    import xlsxwriter  # shipped with Odoo (requirements.txt)

    buffer = io.BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    bold = workbook.add_format({"bold": True, "bg_color": "#D9EEEF", "border": 1})
    wrap = workbook.add_format({"text_wrap": True, "valign": "top"})
    for title, header, rows, widths in sheets:
        sheet = workbook.add_worksheet(title)
        for column, name in enumerate(header):
            sheet.write(0, column, name, bold)
            sheet.set_column(column, column, widths[column] if column < len(widths) else 20, wrap)
        for row_index, row in enumerate(rows, start=1):
            for column, value in enumerate(row):
                if value not in (None, ""):
                    sheet.write(row_index, column, value)
        if header:
            sheet.freeze_panes(1, 0)
    workbook.close()
    return buffer.getvalue()
